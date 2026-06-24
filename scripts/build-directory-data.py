from __future__ import annotations

import json
import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "SMAEndpointDirectory.xlsx"
OUTPUT_PATH = ROOT / "assets" / "directory-data.js"
SHEET_NAME = "SMA Endpoint Directory "

API_GROUPS = [
    {
        "apiType": "Patient Access",
        "start": 6,
        "fields": [
            "productionUrl",
            "status",
            "implementationDate",
            "expectedImplementationDate",
            "vendor",
            "capabilityStatementUrl",
            "registrationInfo",
            "sandboxUrl",
            "developerSupport",
        ],
    },
    {
        "apiType": "Provider Directory API",
        "start": 15,
        "fields": [
            "productionUrl",
            "status",
            "implementationDate",
            "expectedImplementationDate",
            "vendor",
            "capabilityStatementUrl",
            "sandboxUrl",
            "registrationInfo",
            "developerSupport",
        ],
    },
    {
        "apiType": "Provider Access",
        "start": 24,
        "fields": [
            "productionUrl",
            "status",
            "implementationDate",
            "expectedImplementationDate",
            "vendor",
            "capabilityStatementUrl",
            "sandboxUrl",
            "registrationInfo",
            "developerSupport",
        ],
    },
    {
        "apiType": "Payer-to-Payer",
        "start": 33,
        "fields": [
            "productionUrl",
            "status",
            "implementationDate",
            "expectedImplementationDate",
            "vendor",
            "capabilityStatementUrl",
            "sandboxUrl",
            "registrationInfo",
            "developerSupport",
        ],
    },
    {
        "apiType": "Prior Authorization",
        "start": 42,
        "fields": [
            "productionUrl",
            "status",
            "implementationDate",
            "expectedImplementationDate",
            "vendor",
            "capabilityStatementUrl",
            "sandboxUrl",
            "registrationInfo",
            "developerSupport",
        ],
    },
]

UNAVAILABLE_VALUES = {
    "",
    "n/a",
    "na",
    "none",
    "not available",
    "not yet available",
    "not applicable",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_status(value: Any) -> str:
    text = clean(value).lower()
    if not text:
        return "Unknown"

    status_map = {
        "active": "Active",
        "offline": "Offline",
        "in development": "In Development",
        "not yet started": "Not Yet Started",
        "not started": "Not Yet Started",
        "not yet available": "Not Yet Available",
        "n/a": "Not Available",
        "na": "Not Available",
    }
    if text in status_map:
        return status_map[text]
    return " ".join(part.capitalize() for part in text.split(" "))


def is_available_url(value: Any) -> bool:
    text = clean(value)
    return text.lower() not in UNAVAILABLE_VALUES and bool(re.search(r"https?://", text, re.I))


def split_urls(value: Any) -> list[str]:
    text = clean(value)
    if not text:
        return []
    return re.findall(r"https?://[^\s,]+", text)


def cell(row: tuple[Any, ...], one_based_index: int) -> Any:
    zero_based = one_based_index - 1
    if zero_based >= len(row):
        return ""
    return row[zero_based]


def get_directory_sheet(workbook: Any) -> Any:
    if SHEET_NAME in workbook.sheetnames:
        return workbook[SHEET_NAME]

    normalized_sheet_name = SHEET_NAME.strip()
    for sheet_name in workbook.sheetnames:
        if clean(sheet_name) == normalized_sheet_name:
            return workbook[sheet_name]

    return workbook[SHEET_NAME]


def build_data() -> dict[str, Any]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = get_directory_sheet(workbook)
    rows = list(sheet.iter_rows(min_row=3, values_only=True))

    states: list[dict[str, Any]] = []
    endpoints: list[dict[str, Any]] = []

    for row_number, row in enumerate(rows, start=3):
        state = clean(cell(row, 1))
        if not state:
            continue

        state_record = {
            "id": re.sub(r"[^a-z0-9]+", "-", state.lower()).strip("-"),
            "state": state,
            "notes": clean(cell(row, 2)),
            "informationAsOfDate": clean(cell(row, 3)),
            "priorAuthorizationDecisionTimeframesUrl": clean(cell(row, 4)),
            "providerDirectoryWebsiteUrl": clean(cell(row, 5)),
            "providerDirectoryWebsiteAvailable": is_available_url(cell(row, 5)),
            "sourceRow": row_number,
        }
        states.append(state_record)

        for group in API_GROUPS:
            record = {
                "id": f"{state_record['id']}-{group['apiType'].lower().replace(' ', '-').replace('/', '-')}",
                "state": state,
                "apiType": group["apiType"],
                "sourceRow": row_number,
                "informationAsOfDate": state_record["informationAsOfDate"],
            }

            for offset, field in enumerate(group["fields"]):
                record[field] = clean(cell(row, group["start"] + offset))

            record["status"] = normalize_status(record.get("status"))
            record["productionUrls"] = split_urls(record.get("productionUrl"))
            record["capabilityStatementUrls"] = split_urls(record.get("capabilityStatementUrl"))
            record["sandboxUrls"] = split_urls(record.get("sandboxUrl"))
            record["hasProductionUrl"] = bool(record["productionUrls"])
            endpoints.append(record)

    active_count = sum(1 for endpoint in endpoints if endpoint["status"] == "Active")
    available_provider_directories = sum(1 for state in states if state["providerDirectoryWebsiteAvailable"])

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "workbook": WORKBOOK_PATH.name,
        "sourceSheet": sheet.title,
        "states": states,
        "endpoints": endpoints,
        "summary": {
            "stateCount": len(states),
            "endpointCount": len(endpoints),
            "activeEndpointCount": active_count,
            "availableProviderDirectoryCount": available_provider_directories,
            "apiTypes": [group["apiType"] for group in API_GROUPS],
        },
    }


def main() -> None:
    data = build_data()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    js = "window.SMA_DIRECTORY_DATA = "
    js += json.dumps(data, ensure_ascii=False, indent=2)
    js += ";\n"
    OUTPUT_PATH.write_text(js, encoding="utf-8")
    print(
        f"Wrote {OUTPUT_PATH.relative_to(ROOT)} "
        f"with {len(data['states'])} states and {len(data['endpoints'])} endpoint records."
    )


if __name__ == "__main__":
    main()
